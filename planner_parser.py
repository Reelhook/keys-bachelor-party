#!/usr/bin/env python3
"""
planner_parser.py
Parses Key_West_Bachelor_Party_Planner.xlsx and extracts structured JSON for the mobile companion app.
Handles caching based on file modification timestamp.
"""

import os
import re
from datetime import datetime, date
import openpyxl

import urllib.request

APP_DIR = os.path.dirname(os.path.abspath(__file__))
BUNDLED_XLSX = os.path.join(APP_DIR, "Key_West_Bachelor_Party_Planner.xlsx")
GDRIVE_LOCAL_XLSX = "/home/reelsurface/.mnt/gdrive/OldPhoneStuff/Key_West_Bachelor_Party_Planner.xlsx"

DEFAULT_SPREADSHEET_PATH = os.environ.get("SPREADSHEET_PATH") or (GDRIVE_LOCAL_XLSX if os.path.exists(GDRIVE_LOCAL_XLSX) else BUNDLED_XLSX)

def resolve_spreadsheet_path():
    """Returns the best available spreadsheet path or downloads from remote URL if configured."""
    remote_url = os.environ.get("SPREADSHEET_URL")
    sheet_id = os.environ.get("GOOGLE_SHEET_ID")
    drive_id = os.environ.get("GOOGLE_DRIVE_FILE_ID")

    if not remote_url and sheet_id:
        remote_url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    elif not remote_url and drive_id:
        remote_url = f"https://drive.google.com/uc?export=download&id={drive_id}"

    if remote_url:
        try:
            target_tmp = "/tmp/cloud_planner.xlsx"
            req = urllib.request.Request(remote_url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=10) as response, open(target_tmp, "wb") as out_file:
                out_file.write(response.read())
            if os.path.exists(target_tmp) and os.path.getsize(target_tmp) > 1000:
                return target_tmp
        except Exception as e:
            print(f"[Warning] Failed to fetch remote spreadsheet: {e}. Using local fallback.")

    if os.path.exists(GDRIVE_LOCAL_XLSX):
        return GDRIVE_LOCAL_XLSX
    elif os.path.exists(BUNDLED_XLSX):
        return BUNDLED_XLSX
    return DEFAULT_SPREADSHEET_PATH

# Curated coordinates & metadata for Florida Keys venues
VENUE_COORDINATES = {
    "coral lagoon": {"lat": 24.7364, "lng": -81.0118, "category": "lodging", "label": "Coral Lagoon Resort & Marina (MM 53.5)", "icon": "fa-house"},
    "12399 overseas": {"lat": 24.7364, "lng": -81.0118, "category": "lodging", "label": "Coral Lagoon Resort & Marina (MM 53.5)", "icon": "fa-house"},
    "check-out": {"lat": 24.7364, "lng": -81.0118, "category": "lodging", "label": "Coral Lagoon Resort & Marina (MM 53.5)", "icon": "fa-house"},
    "island fish": {"lat": 24.7388, "lng": -81.0142, "category": "food_drink", "label": "Island Fish Co. (MM 54)", "icon": "fa-martini-glass"},
    "havana jack": {"lat": 24.7214, "lng": -81.0125, "category": "food_drink", "label": "Havana Jack's Oceanside", "icon": "fa-utensils"},
    "brass monkey": {"lat": 24.7170, "lng": -81.0664, "category": "nightlife", "label": "The Brass Monkey Lounge (MM 50)", "icon": "fa-beer-mug-empty"},
    "publix": {"lat": 24.7185, "lng": -81.0635, "category": "shopping", "label": "Publix / Winn-Dixie Supermarket", "icon": "fa-cart-shopping"},
    "valhalla": {"lat": 24.7570, "lng": -80.9780, "category": "boating", "label": "Valhalla Sandbar (Crawl Key)", "icon": "fa-anchor"},
    "sombrero reef": {"lat": 24.6275, "lng": -81.1105, "category": "boating", "label": "Sombrero Reef & Sanctuary Light", "icon": "fa-fish"},
    "sunset grille": {"lat": 24.7042, "lng": -81.1278, "category": "food_drink", "label": "Sunset Grille & Raw Bar (7-Mile Bridge)", "icon": "fa-umbrella-beach"},
    "castaway": {"lat": 24.7112, "lng": -81.0912, "category": "food_drink", "label": "Castaway Waterfront Restaurant", "icon": "fa-fish-fins"},
    "dockside": {"lat": 24.7055, "lng": -81.0855, "category": "nightlife", "label": "Dockside Boot Key Harbor", "icon": "fa-music"},
    "mallory": {"lat": 24.5597, "lng": -81.8074, "category": "sightseeing", "label": "Mallory Square (Key West Sunset)", "icon": "fa-sun"},
    "siboney": {"lat": 24.5518, "lng": -81.7925, "category": "food_drink", "label": "El Siboney Cuban Restaurant", "icon": "fa-bowl-rice"},
    "green parrot": {"lat": 24.5528, "lng": -81.8023, "category": "nightlife", "label": "The Green Parrot Bar", "icon": "fa-guitar"},
    "irish kevin": {"lat": 24.5594, "lng": -81.8048, "category": "nightlife", "label": "Irish Kevin's (Duval St)", "icon": "fa-clover"},
    "rick": {"lat": 24.5587, "lng": -81.8049, "category": "nightlife", "label": "Rick's & Durty Harry's Complex", "icon": "fa-champagne-glasses"},
    "burdine": {"lat": 24.7082, "lng": -81.0883, "category": "food_drink", "label": "Burdines Waterfront", "icon": "fa-burger"},
    "sparky": {"lat": 24.7258, "lng": -81.0264, "category": "food_drink", "label": "Sparky's Landing", "icon": "fa-pizza-slice"},
    "jj's": {"lat": 24.7235, "lng": -81.0450, "category": "nightlife", "label": "JJ's Doghouse Sports Bar", "icon": "fa-bullseye"},
    "keys fisheries": {"lat": 24.7155, "lng": -81.0880, "category": "food_drink", "label": "Keys Fisheries Upstairs Marina Deck", "icon": "fa-shrimp"},
    "fishing charter": {"lat": 24.7160, "lng": -81.0890, "category": "boating", "label": "Deep-Sea Fishing Charter Dock", "icon": "fa-ship"},
    "florida keys brewing": {"lat": 24.9228, "lng": -80.6288, "category": "food_drink", "label": "Florida Keys Brewing Co.", "icon": "fa-beer-mug-empty"},
    "islamorada brewery": {"lat": 24.9272, "lng": -80.6225, "category": "food_drink", "label": "Islamorada Brewery & Distillery", "icon": "fa-bottle-droplet"},
    "lorelei": {"lat": 24.9242, "lng": -80.6280, "category": "food_drink", "label": "Lorelei Restaurant & Cabana Bar", "icon": "fa-umbrella-beach"},
}

def match_coordinates(text_to_search, priority_text=""):
    # First check priority_text (e.g. activity name or specific destination)
    if priority_text:
        lower_p = priority_text.lower()
        for key, geo in VENUE_COORDINATES.items():
            if key in lower_p and key != "12399 overseas":
                res = dict(geo)
                return res

    if not text_to_search:
        return None
    lower = text_to_search.lower()
    for key, geo in VENUE_COORDINATES.items():
        if key in lower:
            res = dict(geo)
            return res
    return None

class PlannerParser:
    def __init__(self, excel_path=None):
        self.excel_path = excel_path or resolve_spreadsheet_path()
        self._cached_data = None
        self._last_mtime = None

    def get_data(self, force_refresh=False):
        # Refresh path in case remote URL or fallback is used
        if os.environ.get("SPREADSHEET_URL") or os.environ.get("GOOGLE_SHEET_ID") or os.environ.get("GOOGLE_DRIVE_FILE_ID"):
            self.excel_path = resolve_spreadsheet_path()

        if not os.path.exists(self.excel_path):
            if os.path.exists(BUNDLED_XLSX):
                self.excel_path = BUNDLED_XLSX
            else:
                raise FileNotFoundError(f"Spreadsheet not found at {self.excel_path}")

        current_mtime = os.path.getmtime(self.excel_path)
        if not force_refresh and self._cached_data and self._last_mtime == current_mtime:
            return self._cached_data

        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        
        parsed = {
            "meta": {
                "trip_title": "Jake's Bachelor Party — Florida Keys",
                "dates": "September 10–14, 2026",
                "location": "Coral Lagoon Resort, Marathon FL (MM 53.5)",
                "crew_count": 5,
                "file_modified_time": datetime.fromtimestamp(current_mtime).isoformat(),
                "last_synced": datetime.now().isoformat()
            },
            "dashboard": self._parse_dashboard(wb),
            "itinerary": self._parse_itinerary(wb),
            "activities": self._parse_activities(wb),
            "budget": self._parse_budget(wb),
            "provisioning": self._parse_provisioning(wb),
            "roster": self._parse_roster(wb)
        }

        self._cached_data = parsed
        self._last_mtime = current_mtime
        return parsed

    def _parse_dashboard(self, wb):
        if "Trip Dashboard" not in wb.sheetnames:
            return {}
        ws = wb["Trip Dashboard"]
        data = {
            "crew_roster": [],
            "budget_summary": [],
            "pending_items": [],
            "contacts": []
        }
        
        section = None
        for r in range(1, 45):
            val_a = ws.cell(r, 1).value
            val_b = ws.cell(r, 2).value
            val_c = ws.cell(r, 3).value
            
            str_a = str(val_a or "").strip()
            str_b = str(val_b or "").strip()
            str_c = str(val_c or "").strip()

            if "CREW ROSTER" in str_a:
                section = "crew"
                continue
            elif "BUDGET SUMMARY" in str_a:
                section = "budget"
                continue
            elif "PENDING ACTION ITEMS" in str_a:
                section = "pending"
                continue
            elif "KEY CONTACTS" in str_a:
                section = "contacts"
                continue

            if not str_a and not str_b:
                continue

            if section == "crew" and str_a and "Crew Member" not in str_a:
                data["crew_roster"].append({
                    "name": str_a,
                    "role": str_b
                })
            elif section == "budget" and str_a and "Metric" not in str_a:
                data["budget_summary"].append({
                    "metric": str_a,
                    "amount": str_b
                })
            elif section == "pending" and str_a and "Action Item" not in str_a:
                data["pending_items"].append({
                    "action": str_a,
                    "contact_details": str_b,
                    "owner": str_c
                })
            elif section == "contacts" and str_a and "Service" not in str_a:
                phone_num = re.findall(r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', str_c)
                data["contacts"].append({
                    "service": str_a,
                    "name": str_b,
                    "phone": str_c,
                    "dial_number": phone_num[0] if phone_num else ""
                })

        return data

    def _parse_itinerary(self, wb):
        if "Itinerary" not in wb.sheetnames:
            return {"days": [], "all_items": []}
        ws = wb["Itinerary"]
        
        days_map = {}
        current_day_header = "General"
        
        for r in range(2, ws.max_row + 1):
            val_a = ws.cell(r, 1).value
            val_c = ws.cell(r, 3).value
            str_a = str(val_a or "").strip()
            str_c = str(val_c or "").strip()

            if not str_a and not str_c:
                continue

            # Header row detector
            if any(day_word in str_a.upper() for day_word in ["THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY", "MONDAY"]):
                current_day_header = str_a
                if current_day_header not in days_map:
                    days_map[current_day_header] = []
                continue

            # Skip table header row
            if "Activity / Location" in str_c or str_a == "Date":
                continue

            raw_date = ws.cell(r, 1).value
            if isinstance(raw_date, (datetime, date)):
                date_str = raw_date.strftime("%Y-%m-%d")
            else:
                date_str = str(raw_date or "").split(" ")[0]

            crew_count = ws.cell(r, 2).value
            activity = str_c
            address = str(ws.cell(r, 4).value or "").strip()
            time_slot = str(ws.cell(r, 5).value or "").strip()
            notes = str(ws.cell(r, 6).value or "").strip()
            cost_person = str(ws.cell(r, 7).value or "").strip()
            cost_group = str(ws.cell(r, 8).value or "").strip()
            status = str(ws.cell(r, 9).value or "").strip()

            geo = match_coordinates(f"{activity} {address}", priority_text=activity)

            item = {
                "id": f"itin_{r}",
                "day_header": current_day_header,
                "date": date_str,
                "crew_count": crew_count,
                "activity": activity,
                "address": address,
                "time": time_slot,
                "notes": notes,
                "cost_per_person": cost_person,
                "cost_group": cost_group,
                "status": status,
                "geo": geo
            }

            if current_day_header not in days_map:
                days_map[current_day_header] = []
            days_map[current_day_header].append(item)

        days_list = []
        all_items = []
        for day_title, items in days_map.items():
            days_list.append({
                "title": day_title,
                "items": items
            })
            all_items.extend(items)

        return {"days": days_list, "all_items": all_items}

    def _parse_activities(self, wb):
        if "Activities" not in wb.sheetnames:
            return []
        ws = wb["Activities"]
        activities = []

        for r in range(2, ws.max_row + 1):
            name = str(ws.cell(r, 1).value or "").strip()
            if not name:
                continue

            category = str(ws.cell(r, 2).value or "").strip()
            address = str(ws.cell(r, 3).value or "").strip()
            phone = str(ws.cell(r, 4).value or "").strip()
            res_policy = str(ws.cell(r, 5).value or "").strip()
            hours = str(ws.cell(r, 6).value or "").strip()
            notes = str(ws.cell(r, 7).value or "").strip()
            link = str(ws.cell(r, 8).value or "").strip()
            status = str(ws.cell(r, 9).value or "").strip()

            phone_match = re.findall(r'\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', phone)

            geo = match_coordinates(f"{name} {address}")

            activities.append({
                "id": f"act_{r}",
                "name": name,
                "category": category,
                "address": address,
                "phone": phone,
                "dial_number": phone_match[0] if phone_match else "",
                "reservation_policy": res_policy,
                "hours": hours,
                "notes": notes,
                "link": link,
                "status": status,
                "geo": geo
            })

        return activities

    def _parse_budget(self, wb):
        if "Budget Tracker" not in wb.sheetnames:
            return {"items": [], "categories": {}, "totals": {}}
        ws = wb["Budget Tracker"]
        items = []
        categories = {}
        total_trip_cost = 0.0
        total_paid = 0.0
        total_due = 0.0

        for r in range(2, ws.max_row + 1):
            cat = str(ws.cell(r, 1).value or "").strip()
            item_name = str(ws.cell(r, 2).value or "").strip()
            if not cat and not item_name:
                continue

            total_val = ws.cell(r, 3).value
            payer = str(ws.cell(r, 4).value or "").strip()
            split_num = ws.cell(r, 5).value
            per_person_val = ws.cell(r, 6).value
            status = str(ws.cell(r, 7).value or "").strip()

            if "TOTAL" in cat.upper() or "TOTAL" in item_name.upper():
                continue

            cost = float(total_val) if isinstance(total_val, (int, float)) else 0.0
            per_person = float(per_person_val) if isinstance(per_person_val, (int, float)) else (cost / 5.0 if cost else 0.0)

            total_trip_cost += cost
            if "PAID" in status.upper():
                total_paid += cost
            else:
                total_due += cost

            if cat not in categories:
                categories[cat] = 0.0
            categories[cat] += cost

            items.append({
                "category": cat,
                "item": item_name,
                "total_cost": round(cost, 2),
                "payer": payer,
                "split_among": split_num,
                "per_person": round(per_person, 2),
                "status": status
            })

        return {
            "items": items,
            "category_breakdown": {k: round(v, 2) for k, v in categories.items()},
            "summary": {
                "total_cost": round(total_trip_cost, 2),
                "per_person_estimate": round(total_trip_cost / 5.0, 2),
                "total_paid": round(total_paid, 2),
                "total_due": round(total_due, 2)
            }
        }

    def _parse_provisioning(self, wb):
        if "Provisioning" not in wb.sheetnames:
            return {"categories": [], "all_items": []}
        ws = wb["Provisioning"]
        
        categories_dict = {}
        all_items = []
        current_category = "General"

        for r in range(2, ws.max_row + 1):
            cat = str(ws.cell(r, 1).value or "").strip()
            desc = str(ws.cell(r, 2).value or "").strip()
            
            if not cat and not desc:
                continue

            if "Subtotal:" in cat or "Subtotal:" in desc or "TOTAL" in cat.upper():
                continue

            if cat:
                current_category = cat

            brand = str(ws.cell(r, 3).value or "").strip()
            qty = ws.cell(r, 4).value
            unit_cost = ws.cell(r, 5).value
            total_cost = ws.cell(r, 6).value
            assigned = str(ws.cell(r, 7).value or "").strip()
            status = str(ws.cell(r, 8).value or "").strip()

            cost = float(total_cost) if isinstance(total_cost, (int, float)) else 0.0

            item = {
                "id": f"prov_{r}",
                "category": current_category,
                "description": desc,
                "brand": brand,
                "qty": qty,
                "unit_cost": unit_cost,
                "total_cost": cost,
                "assigned": assigned,
                "status": status
            }

            if current_category not in categories_dict:
                categories_dict[current_category] = []
            categories_dict[current_category].append(item)
            all_items.append(item)

        categories_list = [{"name": cat_name, "items": items} for cat_name, items in categories_dict.items()]

        return {"categories": categories_list, "all_items": all_items}

    def _parse_roster(self, wb):
        if "Roster Availability" not in wb.sheetnames:
            return []
        ws = wb["Roster Availability"]
        roster = []

        headers = [str(ws.cell(1, c).value or "").strip() for c in range(2, 7)]
        subheaders = [str(ws.cell(3, c).value or "").strip() for c in range(2, 7)]

        for r in range(4, 15):
            name = str(ws.cell(r, 1).value or "").strip()
            if not name:
                continue
            availability = {}
            for c_idx, day_name in enumerate(headers):
                status_val = str(ws.cell(r, c_idx + 2).value or "").strip()
                desc = subheaders[c_idx] if c_idx < len(subheaders) else ""
                availability[day_name] = {
                    "theme": desc,
                    "status": status_val
                }
            roster.append({
                "name": name,
                "availability": availability
            })

        return roster

if __name__ == "__main__":
    parser = PlannerParser()
    data = parser.get_data()
    print("Parsed trip:", data["meta"]["trip_title"])
    print(f"Days: {len(data['itinerary']['days'])}, Activities: {len(data['activities'])}, Budget Items: {len(data['budget']['items'])}, Provisioning items: {len(data['provisioning']['all_items'])}")
